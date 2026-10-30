import os
import re
from typing import Dict, Any

def extract_text_from_pdf(file_path: str) -> str:
    """Extracts text from PDF using pypdf or PyMuPDF (fitz)."""
    # Try pypdf first (pure Python, lightweight, serverless friendly)
    try:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        text_parts = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        text = "\n".join(text_parts).strip()
        if text:
            return text
    except Exception:
        pass

    # Fallback to PyMuPDF if available
    try:
        import fitz
        doc = fitz.open(file_path)
        text_parts = []
        for page in doc:
            text_parts.append(page.get_text("text"))
        text = "\n".join(text_parts).strip()
        return text
    except Exception as e:
        raise ValueError(f"Could not read PDF file: {str(e)}")

def extract_text_from_docx(file_path: str) -> str:
    """Extracts text from DOCX using python-docx."""
    try:
        import docx
        doc = docx.Document(file_path)
        text_parts = [para.text for para in doc.paragraphs if para.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
        return "\n".join(text_parts).strip()
    except Exception as e:
        raise ValueError(f"Could not read DOCX file: {str(e)}")

def extract_text_from_file(file_path: str, filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        text = extract_text_from_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        text = extract_text_from_docx(file_path)
    elif ext in [".txt"]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read().strip()
    else:
        raise ValueError(f"Unsupported file format: {ext}. Please upload a PDF or DOCX file.")
    
    if not text or len(text.strip()) < 40:
        raise ValueError("The uploaded document appears to be empty, scanned image-only, or unreadable. Please upload a text-readable PDF or DOCX.")
    
    return text

def parse_resume_heuristics(text: str) -> Dict[str, Any]:
    """Lightweight deterministic entity extraction for fast metadata."""
    email_match = re.search(r"[\w\.-]+@[\w\.-]+\.\w+", text)
    phone_match = re.search(r"(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}", text)
    
    # Common tech skills dictionary
    common_skills = [
        "Python", "JavaScript", "TypeScript", "React", "Node.js", "Express", "Next.js",
        "HTML", "CSS", "SQL", "PostgreSQL", "MongoDB", "MySQL", "Redis",
        "Docker", "Kubernetes", "AWS", "Azure", "GCP", "CI/CD", "Git", "GitHub",
        "REST API", "GraphQL", "FastAPI", "Flask", "Django", "Java", "Spring Boot",
        "C++", "C#", ".NET", "Go", "Rust", "Tailwind CSS", "Linux",
        "Machine Learning", "Data Analysis", "Pandas", "NumPy", "PyTorch", "TensorFlow",
        "Agile", "Scrum", "Jira", "Figma", "Redux", "Kafka", "Microservices"
    ]
    
    lower_text = text.lower()
    detected_skills = []
    for skill in common_skills:
        # Match as full word
        pattern = r"\b" + re.escape(skill.lower()) + r"\b"
        if re.search(pattern, lower_text):
            detected_skills.append(skill)
            
    # Simple experience estimate
    years_match = re.search(r"(\d+)\+?\s*(?:years?|yrs?)\s*(?:of)?\s*(?:experience|exp)", lower_text)
    years = float(years_match.group(1)) if years_match else 0.0

    return {
        "email": email_match.group(0) if email_match else None,
        "phone": phone_match.group(0) if phone_match else None,
        "detected_skills": detected_skills,
        "estimated_years_experience": years,
        "char_count": len(text),
        "word_count": len(text.split())
    }
